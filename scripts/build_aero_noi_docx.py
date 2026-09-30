"""Build five course content review pages from their shared content file."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
STEM = '2026-09-30-aero-noi-review'
courses = json.loads((ROOT / 'curriculum-content' / (STEM + '.json')).read_text())
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.5)
sec.left_margin = sec.right_margin = Inches(.6)
sec.footer_distance = Inches(.22)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.widow_control = True
for name, size in [('Title', 15), ('Subtitle', 11), ('Heading 1', 10.5)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
h = doc.styles['Heading 1'].paragraph_format
h.space_before = Pt(7)
h.space_after = Pt(3)
h.keep_with_next = True
footer = sec.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run('AERO notice of intent  |  ')
run.font.size = Pt(8)
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)
doc.core_properties.title = 'AERO Notices of Intent'
doc.core_properties.author = 'Kyle Nessen'
doc.core_properties.subject = 'Course proposals for Ron Overacker review'
typ = ['#set document(title: "AERO Notices of Intent", author: "Kyle Nessen")',
       '#set page(paper: "us-letter", margin: (x: 0.6in, y: 0.5in), numbering: "1")',
       '#set text(font: "Arial", size: 10pt)', '#set par(leading: 0.3em, spacing: 0.35em)',
       '#set heading(numbering: none)', '#show heading.where(level: 1): set text(size: 15pt)',
       '#show heading.where(level: 2): set text(size: 10.5pt)']

def para(label, value):
    p = doc.add_paragraph()
    p.add_run(label + '  ').bold = True
    p.add_run(value)
    typ.append('*' + label + '* ' + value.replace('_', r'\_') + '\n')

def heading(s):
    doc.add_paragraph(s, 'Heading 1')
    typ.append('== ' + s + '\n')

def body(value):
    doc.add_paragraph(value)
    typ.append(value.replace('_', r'\_') + '\n')

for i,c in enumerate(courses):
    if i:
        doc.add_page_break()
        typ.append('#pagebreak()')
    n = c['number']
    independent = n == '147'
    structure = c['structure']
    course_heading = 'AERO ' + n + (' Proposed' if independent else '')
    doc.add_paragraph(course_heading, 'Title')
    doc.add_paragraph(c['title'], 'Subtitle')
    typ += ['= ' + course_heading, '*' + c['title'] + '*\n']
    units = '0.5 to 2 lab units' if independent else f"{structure['minimum_units']} units, {structure['lecture_units']} lecture and {structure['lab_units']} lab"
    para('Course structure', units + f". Enrollment cap {structure['cap']}." + (' Small lab.' if structure['cap_category'] else ''))
    para('Semester hours', f"{structure['lecture_hours_per_term']} lecture, {structure['lab_hours_per_term']} lab, {structure['outside_hours_per_term']} outside work, and 0 activity. {structure['total_student_hours']} total student hours.")
    para('Minimum qualifications', 'Aeronautics')
    heading('Course description')
    body(c['description'])
    heading('Course rationale')
    body(c['rationale'])
    heading('Connection to the College Mission')
    body(c['mission'])
    heading('Program fit')
    body('The course does not fit within an existing degree or certificate. ' + ('' if independent else 'It will support proposed certificates of specialization. ') + 'Transfer or articulation is not proposed.')
    heading('Related courses and consultation')
    body(c['overlap'].removeprefix('Yes. '))
    heading('Enrollment and department efficiency')
    body(c['efficiency'])
    heading('Resources and funding')
    body(structure['resources_and_funding'])

out = ROOT / 'output' / 'docx' / (STEM + '.docx')
out.parent.mkdir(parents=True, exist_ok=True)
for element in list(doc.styles.element.iter(qn('w:pBdr'))):
    element.getparent().remove(element)
for element in list(doc.element.iter(qn('w:pBdr'))):
    element.getparent().remove(element)
doc.save(out)
(ROOT / 'curriculum-content' / (STEM + '.typ')).write_text('\n'.join(typ)+'\n')
print(out)
